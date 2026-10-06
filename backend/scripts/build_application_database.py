"""Build the application database from the market survey Excel and TDS PDFs.

Run with:  python -m scripts.build_application_database   (from backend/)

Inputs (all under seed/data/):
  source/Graphene Application Database.xlsx   hand-compiled vendor survey
  source/GRMs_Tables_S1_S2.xlsx               MDPI Carbon 2026 supplementary
                                              tables, ~390 commercial GRMs
  tds_pdfs/*.pdf                              vendor technical data sheets

Outputs:
  seed/data/application_products.json   one record per vendor grade, with
                                        every spec as a {low, high, text}
                                        interval and canonical application
                                        tags
  seed/data/market_reference.json       numeric spec intervals for the wider
                                        market, used for percentile context

Two things make this more than a CSV export. Excel silently converted every
"m-d"-shaped range ("2-3" layers, "5-25" um) to a date; those are recovered
from the cell's number format. And the PDFs were read against the sheet:
where they disagree the datasheet wins, and each override is listed in
CORRECTIONS with its reason so the provenance is auditable.
"""

from __future__ import annotations

import datetime as dt
import json
import re
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.analysis.specvalue import SpecValue, parse_spec_value  # noqa: E402

DATA_DIR = Path(__file__).resolve().parents[1] / "seed" / "data"
SURVEY_XLSX = DATA_DIR / "source" / "Graphene Application Database.xlsx"
MARKET_XLSX = DATA_DIR / "source" / "GRMs_Tables_S1_S2.xlsx"
PDF_DIR = DATA_DIR / "tds_pdfs"
TAXONOMY = DATA_DIR / "application_taxonomy.json"
PRODUCTS_OUT = DATA_DIR / "application_products.json"
MARKET_OUT = DATA_DIR / "market_reference.json"

# Survey sheet columns (0-based) -> canonical spec key and unit scale.
# Column 8 is "Number of Layers", column 15 electrical conductivity in S/m.
SPEC_COLUMNS: dict[int, tuple[str, float]] = {
    7: ("lateral_size_um", 1.0),
    8: ("layers", 1.0),
    9: ("bet_m2_g", 1.0),
    11: ("carbon_purity_pct", 1.0),
    12: ("oxygen_pct", 1.0),
    13: ("impurities_pct", 1.0),
    14: ("bulk_density_g_cm3", 1.0),
    15: ("electrical_conductivity_s_m", 1.0),
    16: ("thermal_conductivity_w_mk", 1.0),
    17: ("id_ig", 1.0),
    18: ("loading_wt_pct", 1.0),
    19: ("density_g_ml", 1.0),
    20: ("viscosity_cps", 1.0),
    21: ("sheet_resistance_ohm_sq", 1.0),
}
TEXT_COLUMNS: dict[int, str] = {
    10: "surface_chemistry",
    22: "solvent",
    23: "additives",
    24: "notes",
}

# Overrides applied after parsing, keyed by product name. Each entry replaces
# a spec interval (or text field) and says why, so nothing is silently edited.
CORRECTIONS: dict[str, list[tuple[str, object, str]]] = {
    "RGA-COOH-1": [
        ("lateral_size_um", "0.02-0.05", "Datasheet: lateral dimensions 20-50 nm by TEM; sheet said 20-50 um"),
        ("bulk_density_g_cm3", "0.2-0.27", "Datasheet: tapped bulk density 200-270 mg/mL"),
        ("layers", "3-9", "Datasheet: 6-layer average, 3-9 layers"),
    ],
    "FGP-2000-AQ-1": [
        ("lateral_size_um", "0.02-0.05", "Datasheet: primary lateral size 20-50 nm (TEM), D50 35 nm"),
        ("layers", "3-9", "Datasheet: thickness 3-9 layers"),
        ("loading_wt_pct", "20", "Datasheet: 20 wt% FGA-1 graphene solids; sheet gave the fraction 0.2"),
        ("viscosity_cps", "14-51", "Datasheet: 0.051 Pa.s at 100 1/s, 0.014 Pa.s at 1000 1/s"),
    ],
    "FGA-1": [
        ("bulk_density_g_cm3", "0.07-0.1", "Datasheet: 70-100 mg/mL manual tapping"),
        ("electrical_conductivity_s_m", "100-300", "Datasheet: 100-300 S/m, function of powder compression"),
        ("layers", "3-10", "Datasheet: 6-layer average, <=10 layers"),
    ],
    "FGA-2": [
        ("bulk_density_g_cm3", "0.07-0.1", "Datasheet: 70-100 mg/mL manual tapping"),
    ],
    "ADG-A+": [
        ("bulk_density_g_cm3", "0.014", "Datasheet: 0.014 g/cm3 (fluffy powder); sheet said 0.14"),
    ],
    "PureGraph AQUA": [
        ("lateral_size_um", "5-50", "Datasheet: grades with D50 of 5, 10, 20 and 50 um"),
        ("loading_wt_pct", "20", "Datasheet: 20% w/w graphene in water"),
    ],
    "PureGraph MB-LDPE": [
        ("lateral_size_um", "5-50", "Datasheet: grades with DV50 of 5, 10, 20 and 50 um"),
    ],
    "PureGraph MB-EVA": [
        ("lateral_size_um", "5-50", "Datasheet: grades with DV50 of 5, 10, 20 and 50 um"),
        ("layers", ">15", "Sheet said >16; kept consistent with the sibling PureGRAPH grades"),
    ],
    "Graphene Nanosheets": [
        ("lateral_size_um", "0.45", "Datasheet: average size ~450 nm (DLS)"),
        ("oxygen_pct", "1", "Datasheet: oxygen ~1%"),
    ],
    "xGnP Graphene Nanoplatelets - Grade H": [
        ("thermal_conductivity_w_mk", "6-3000", "Sheet cell was date-corrupted: 3000 W/mK in-plane, 6 W/mK through-plane"),
    ],
    "G3HT": [
        ("bulk_density_g_cm3", "0.06", "Sheet value 60 is only physical as mg/mL (0.06 g/cm3)"),
    ],
    "G1HT": [
        ("bulk_density_g_cm3", "0.03", "Sheet value 30 is only physical as mg/mL (0.03 g/cm3)"),
    ],
    "Industrial Grade Graphene Nanoplatelets": [
        ("lateral_size_um", "0.5-1", "Product page title: 6-10 layer, 0.5-1 um; the sheet had the two columns swapped"),
        ("layers", "6-10", "Product page title: 6-10 layer, 0.5-1 um; the sheet had the two columns swapped"),
    ],
    "GRA-350": [
        ("lateral_size_um", "0.1-0.8", "Sheet typo '01-0.8'"),
    ],
    "GRA-357": [
        ("bulk_density_g_cm3", None, "Sheet value 7.7 g/cm3 exceeds graphite density (2.26); dropped as unphysical"),
    ],
}

# Vendor application text that the keyword mapper cannot resolve on its own.
EXTRA_TAGS: dict[str, list[str]] = {
    "G1HT": ["polymer_composites", "concrete_cement", "coatings_paints", "batteries"],  # older grade of G3HT
    # Read off the datasheet application panels, which list more than the sheet.
    "ADG-A+": ["sensors_electronics"],
    "ADG-X": ["elastomers_rubber_tyre", "sensors_electronics"],
    "ADG-D": ["sensors_electronics"],
}

# The sheet row for these products carried text copied from another vendor;
# the datasheet's own application sentence is used instead.
APPLICATION_TEXT_OVERRIDES: dict[str, str] = {
    "Graphene Nanosheets": (
        "Energy storage (batteries, supercapacitors), electronics, electromagnetic "
        "shielding, sensors, catalysis, advanced composite materials"
    ),
}

FORM_ALIASES = {
    "powder": "powder",
    "dispersion": "dispersion",
    "paste": "paste",
    "pellet": "pellet",
    "masterbatch": "pellet",
    "slurry": "dispersion",
    "liquid": "dispersion",
}

COMPANY_WEBSITES = {
    "Techinstro": "https://www.techinstro.com/",
    "Shipent": "https://www.shilpent.com/",
    "First Graphene": "https://firstgraphene.net/",
    "Dycotec Material": "https://www.dycotecmaterials.com/",
    "Hydrograph": "https://hydrograph.com/",
    "The Sixth Element Inc.": "https://www.c6th.com/",
    "GraphenEra": "https://graphenera.com/",
    "NanoXplore": "https://nanoxplore.ca/",
    "Levidian": "https://levidian.com/",
    "GrapheneUP": "https://grapheneup.com/",
    "AdNano Technologies Pvt. Ltd.": "https://www.ad-nanotech.com/",
    "Cheap Tubes Inc.": "https://www.cheaptubes.com/",
    "NanoGrafi": "https://nanografi.com/",
    "Matexcel": "https://www.matexcel.com/",
    "GTechPlasma": "https://gtechplasma.com/",
}


def _cell_text(cell) -> str | None:
    """Recover ranges that Excel turned into dates; otherwise return the text."""
    value = cell.value
    if value is None:
        return None
    if isinstance(value, (dt.datetime, dt.date)):
        fmt = (cell.number_format or "").lower()
        if "yyyy" in fmt or "yy" in fmt:
            # "6-3000" -> 3000-06-01 with format m-yyyy
            return f"{value.month}-{value.year}"
        # "2-3" -> 2026-02-03 with format m-d
        return f"{value.month}-{value.day}"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return text or None


def _slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", text.lower()).strip("_")


def _split_applications(text: str | None) -> list[str]:
    if not text or text.strip() in {"-", ""}:
        return []
    parts = re.split(r"[,;/]|\band\b|\bsuch as\b|\betc\b|\.", text)
    return [p.strip(" .") for p in parts if p.strip(" .")]


def map_applications(text: str | None, taxonomy: list[dict]) -> list[str]:
    tags: list[str] = []
    fragments = _split_applications(text)
    haystack = " ".join(fragments).lower()
    for entry in taxonomy:
        for keyword in entry["keywords"]:
            pattern = r"(?<![a-z])" + re.escape(keyword.lower()) + r"(?![a-z])"
            if re.search(pattern, haystack):
                tags.append(entry["key"])
                break
    return tags


def pdf_text(path: Path) -> str:
    """Extract text once and drop repeated lines (many TDS repeat the front page)."""
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    seen: set[str] = set()
    lines: list[str] = []
    for page in reader.pages:
        for line in (page.extract_text() or "").splitlines():
            clean = re.sub(r"\s+", " ", line).strip()
            if not clean or clean in seen:
                continue
            seen.add(clean)
            lines.append(clean)
    return "\n".join(lines)


def _match_pdf(link: str | None, product: str, pdfs: dict[str, Path]) -> Path | None:
    if link:
        name = Path(link).name
        if name in pdfs:
            return pdfs[name]
        stem = Path(name).stem.lower()
        for filename, path in pdfs.items():
            if stem and stem in filename.lower():
                return path
    for filename, path in pdfs.items():
        if product.lower().replace(" ", "") in filename.lower().replace(" ", ""):
            return path
    return None


def build_products(taxonomy: list[dict]) -> list[dict]:
    wb = openpyxl.load_workbook(SURVEY_XLSX, data_only=True)
    ws = wb["Sheet1"]
    pdfs = {p.name: p for p in PDF_DIR.glob("*.pdf")} if PDF_DIR.is_dir() else {}

    products: list[dict] = []
    for row in ws.iter_rows(min_row=3):
        cells = list(row)
        company = _cell_text(cells[1])
        product = _cell_text(cells[2])
        if not company or not product:
            continue
        company = company.strip()
        product = product.strip()

        form_text = (_cell_text(cells[3]) or "").lower()
        form = next((canon for alias, canon in FORM_ALIASES.items() if alias in form_text), "unknown")

        specs: dict[str, dict] = {}
        for column, (key, scale) in SPEC_COLUMNS.items():
            parsed = parse_spec_value(_cell_text(cells[column]), scale)
            if parsed is not None:
                specs[key] = parsed.to_dict()

        text_fields = {}
        for column, key in TEXT_COLUMNS.items():
            value = _cell_text(cells[column])
            if value and value not in {"-", "–"}:
                text_fields[key] = value

        # For liquids the "bulk density" column holds the dispersion density.
        if form in {"dispersion", "paste"} and "bulk_density_g_cm3" in specs and "density_g_ml" not in specs:
            specs["density_g_ml"] = specs.pop("bulk_density_g_cm3")

        corrections_applied: list[dict] = []
        for key, replacement, reason in CORRECTIONS.get(product, []):
            if replacement is None:
                specs.pop(key, None)
            else:
                parsed = parse_spec_value(replacement)
                if parsed is not None:
                    specs[key] = parsed.to_dict()
            corrections_applied.append({"field": key, "value": replacement, "reason": reason})

        applications_text = APPLICATION_TEXT_OVERRIDES.get(product, _cell_text(cells[5]))
        tags = map_applications(applications_text, taxonomy)
        for tag in EXTRA_TAGS.get(product, []):
            if tag not in tags:
                tags.append(tag)

        link = _cell_text(cells[6])
        pdf_path = _match_pdf(link, product, pdfs)
        datasheet_url = link if link and link.startswith("http") else None
        if datasheet_url is None:
            datasheet_url = COMPANY_WEBSITES.get(company)

        record = {
            "key": f"{_slug(company)}__{_slug(product)}",
            "company": company,
            "product": product,
            "form": form,
            "synthesis": (_cell_text(cells[4]) or "").strip(" -") or None,
            "applications_text": applications_text,
            "application_tags": tags,
            "specs": specs,
            **text_fields,
            "datasheet_url": datasheet_url,
            "datasheet_file": pdf_path.name if pdf_path else None,
            "datasheet_text": pdf_text(pdf_path) if pdf_path else None,
            "corrections": corrections_applied,
            "source_row": cells[0].row,
        }
        products.append(record)
    return products


# --------------------------------------------------------------------------- #
# Market reference (MDPI Carbon 2026 supplementary tables)
# --------------------------------------------------------------------------- #
MARKET_COLUMNS: dict[str, tuple[str, str]] = {
    "Layers": ("layers", ""),
    "Thickness": ("thickness_nm", "nm"),
    "Specific Surface Area": ("bet_m2_g", "m2/g"),
    "Lateral Dimension": ("lateral_size_um", "um"),
    "D/G Intensity Peak ratio": ("id_ig", ""),
    "Electrical conductivity": ("electrical_conductivity_s_m", "S/m"),
    "Elemental C": ("carbon_purity_pct", "%"),
    "Elemental O": ("oxygen_pct", "%"),
    "Apparent Density": ("bulk_density_g_cm3", "g/cm3"),
}


def _market_value(raw: object, key: str) -> SpecValue | None:
    """Best-effort parse of the literature table cells, which carry units inline."""
    if raw is None:
        return None
    text = str(raw).strip()
    if not text or text.startswith("#") or text.lower() in {"nan", "um", "/ l", "/ l"}:
        return None
    lowered = text.lower()
    scale = 1.0
    if key == "electrical_conductivity_s_m":
        if "s/cm" in lowered.replace(" ", ""):
            scale = 100.0
        if "ω" in lowered or "ohm" in lowered or "insulat" in lowered:
            return None
    if key == "thickness_nm" and "um" in lowered and "nm" not in lowered:
        scale = 1000.0
    if key == "lateral_size_um" and "nm" in lowered:
        scale = 1e-3
    if key == "bulk_density_g_cm3" and "mg" in lowered:
        scale = 1e-3
    # Strip unit words that would otherwise be read as numbers ("g/cm3", "m2/g").
    stripped = re.sub(r"(g\s*/\s*c\s*m\s*[34]|g\s*c\s*m-3|m\s*2\s*/\s*g(?:ram)?|s\s*/\s*c?m|w/mk|mg\s*/\s*ml|g\s*/\s*ml|wt%|at%|nm|um|µm|%|c:)", " ", lowered)
    stripped = stripped.replace("<0.005", "<0.005")
    parsed = parse_spec_value(stripped, scale)
    if parsed is None:
        return None
    # Sanity windows keep unit-garbled cells out of the percentiles.
    windows = {
        "layers": (1, 200),
        "thickness_nm": (0.3, 200),
        "bet_m2_g": (1, 3000),
        "lateral_size_um": (0.001, 500),
        "id_ig": (0.001, 3),
        "electrical_conductivity_s_m": (1e-3, 1e8),
        "carbon_purity_pct": (20, 100),
        "oxygen_pct": (0, 60),
        "bulk_density_g_cm3": (0.001, 2.3),
    }
    low, high = windows[key]
    for bound in (parsed.low, parsed.high):
        if bound is not None and not (low <= bound <= high):
            return None
    return SpecValue(parsed.low, parsed.high, text)


def build_market_reference() -> list[dict]:
    if not MARKET_XLSX.is_file():
        return []
    wb = openpyxl.load_workbook(MARKET_XLSX, data_only=True, read_only=True)
    s1 = list(wb["Table S1"].iter_rows(values_only=True))
    s2 = list(wb["Table S2"].iter_rows(values_only=True))
    s1_header = [str(h).strip() if h is not None else "" for h in s1[0]]
    s2_header = [str(h).strip() if h is not None else "" for h in s2[0]]

    def col(header: list[str], name: str) -> int | None:
        return header.index(name) if name in header else None

    s1_company, s1_product = col(s1_header, "Company name and website"), col(s1_header, "Product (GRMs)")
    s1_form, s1_synth, s1_country = col(s1_header, "Product form"), col(s1_header, "Synthesis Process"), col(s1_header, "Country")
    forms: dict[tuple[str, str], dict] = {}
    for row in s1[1:]:
        if s1_company is None or s1_product is None:
            break
        company, product = row[s1_company], row[s1_product]
        if company is None or product is None:
            continue
        forms.setdefault(
            (_slug(re.split(r"\s+https?://", str(company))[0]), str(product).strip().lower()),
            {
                "form": str(row[s1_form]).strip() if s1_form is not None and row[s1_form] else None,
                "synthesis": str(row[s1_synth]).strip() if s1_synth is not None and row[s1_synth] else None,
                "country": str(row[s1_country]).strip() if s1_country is not None and row[s1_country] else None,
            },
        )

    records: list[dict] = []
    company_col, product_col = col(s2_header, "Company name and website"), col(s2_header, "Product (GRMs)")
    for index, row in enumerate(s2[1:], start=1):
        if company_col is None or product_col is None:
            break
        company_raw, product_raw = row[company_col], row[product_col]
        if company_raw is None:
            continue
        company_text = str(company_raw).strip()
        company_name = re.split(r"\s+https?://", company_text)[0].strip()
        url = None
        url_match = re.search(r"https?://\S+", company_text)
        if url_match:
            url = url_match.group(0)
        product = str(product_raw).strip() if product_raw is not None else "Graphene"
        specs: dict[str, dict] = {}
        for column_name, (key, _unit) in MARKET_COLUMNS.items():
            index_in_row = col(s2_header, column_name)
            if index_in_row is None:
                continue
            value = _market_value(row[index_in_row], key)
            if value is not None:
                specs[key] = value.to_dict()
        if not specs:
            continue
        meta = forms.get((_slug(company_name), product.lower()), {})
        records.append(
            {
                "key": f"market__{index:03d}__{_slug(company_name)}",
                "company": company_name,
                "product": product,
                "url": url,
                "form": (meta.get("form") or "").lower().replace(" ", "") or None,
                "synthesis": meta.get("synthesis"),
                "country": meta.get("country"),
                "specs": specs,
            }
        )
    return records


def main() -> int:
    taxonomy = json.loads(TAXONOMY.read_text(encoding="utf-8"))
    products = build_products(taxonomy)
    PRODUCTS_OUT.write_text(json.dumps(products, indent=1, ensure_ascii=False), encoding="utf-8")
    untagged = [p["product"] for p in products if not p["application_tags"]]
    with_pdf = sum(1 for p in products if p["datasheet_file"])
    print(f"application_products.json: {len(products)} products, {with_pdf} with a datasheet PDF")
    if untagged:
        print(f"  no application tag: {', '.join(untagged)}")
    counts: dict[str, int] = {}
    for record in products:
        for tag in record["application_tags"]:
            counts[tag] = counts.get(tag, 0) + 1
    for tag, count in sorted(counts.items(), key=lambda item: -item[1]):
        print(f"  {tag:24s} {count:2d} products")

    market = build_market_reference()
    MARKET_OUT.write_text(json.dumps(market, indent=1, ensure_ascii=False), encoding="utf-8")
    coverage: dict[str, int] = {}
    for record in market:
        for key in record["specs"]:
            coverage[key] = coverage.get(key, 0) + 1
    print(f"market_reference.json: {len(market)} products with at least one numeric spec")
    for key, count in sorted(coverage.items(), key=lambda item: -item[1]):
        print(f"  {key:28s} {count:3d}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
